from __future__ import annotations

import json
import re
import unicodedata
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from .errors import SpreadsheetValidationError
from .hashing import hash_canonical
from .models import NormalizedDocumentRow, WorkbookData


HEADER_ALIASES = {
    "tipo de documento": "tipo_documento",
    "cufe cude": "clave_documento",
    "folio": "folio",
    "prefijo": "prefijo",
    "divisa": "divisa",
    "forma de pago": "forma_pago",
    "medio de pago": "medio_pago",
    "fecha emision": "fecha_emision",
    "fecha emisin": "fecha_emision",
    "fecha recepcion": "fecha_recepcion",
    "fecha recepcin": "fecha_recepcion",
    "nit emisor": "nit_emisor",
    "nombre emisor": "nombre_emisor",
    "nit receptor": "nit_receptor",
    "nombre receptor": "nombre_receptor",
    "iva": "iva",
    "ica": "ica",
    "ic": "ic",
    "inc": "inc",
    "timbre": "timbre",
    "inc bolsas": "inc_bolsas",
    "in carbono": "in_carbono",
    "in combustibles": "in_combustibles",
    "ic datos": "ic_datos",
    "icl": "icl",
    "inpp": "inpp",
    "ibua": "ibua",
    "icui": "icui",
    "rete iva": "rete_iva",
    "rete renta": "rete_renta",
    "rete ica": "rete_ica",
    "total": "total",
    "estado": "estado_dian",
    "grupo": "grupo",
}

REQUIRED_COLUMNS = {
    "tipo_documento",
    "clave_documento",
    "folio",
    "fecha_emision",
    "fecha_recepcion",
    "nit_emisor",
    "nombre_emisor",
    "nit_receptor",
    "nombre_receptor",
    "total",
    "estado_dian",
    "grupo",
}

DECIMAL_COLUMNS = (
    "iva",
    "ica",
    "ic",
    "inc",
    "timbre",
    "inc_bolsas",
    "in_carbono",
    "in_combustibles",
    "ic_datos",
    "icl",
    "inpp",
    "ibua",
    "icui",
    "rete_iva",
    "rete_renta",
    "rete_ica",
)


def read_dian_workbook(
    content: bytes,
    preferred_table_name: str | None = None,
) -> WorkbookData:
    if not content:
        raise SpreadsheetValidationError("El archivo XLSX está vacío.")

    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries

        workbook = load_workbook(BytesIO(content), read_only=False, data_only=True)
    except Exception as error:
        raise SpreadsheetValidationError(
            "No fue posible abrir el archivo como un libro XLSX válido."
        ) from error

    try:
        worksheet, table_name, table_ref = _select_table(
            workbook, preferred_table_name
        )
        min_column, min_row, max_column, max_row = range_boundaries(table_ref)
        raw_headers = [
            worksheet.cell(row=min_row, column=column).value
            for column in range(min_column, max_column + 1)
        ]
        header_keys = _map_headers(raw_headers)

        rows: list[NormalizedDocumentRow] = []
        for row_number in range(min_row + 1, max_row + 1):
            raw_values = [
                worksheet.cell(row=row_number, column=column).value
                for column in range(min_column, max_column + 1)
            ]
            if all(_is_blank(value) for value in raw_values):
                continue
            rows.append(
                _normalize_row(
                    row_number=row_number,
                    raw_headers=raw_headers,
                    header_keys=header_keys,
                    raw_values=raw_values,
                )
            )

        return WorkbookData(
            sheet_name=worksheet.title,
            table_name=table_name,
            rows=tuple(rows),
        )
    finally:
        workbook.close()


def _select_table(workbook, preferred_table_name: str | None):
    tables = []
    for worksheet in workbook.worksheets:
        for table in worksheet.tables.values():
            tables.append((worksheet, table.name, table.ref))

    if preferred_table_name:
        expected = preferred_table_name.strip().casefold()
        for worksheet, name, ref in tables:
            if name.casefold() == expected:
                return worksheet, name, ref
        raise SpreadsheetValidationError(
            f"No existe la tabla solicitada: {preferred_table_name}."
        )

    if tables:
        return tables[0]

    for worksheet in workbook.worksheets:
        if worksheet.max_row >= 2 and worksheet.max_column >= 1:
            return (
                worksheet,
                worksheet.title,
                f"A1:{worksheet.cell(worksheet.max_row, worksheet.max_column).coordinate}",
            )

    raise SpreadsheetValidationError(
        "El libro no contiene una tabla ni una hoja con filas de datos."
    )


def _map_headers(raw_headers: list[Any]) -> list[str | None]:
    mapped: list[str | None] = []
    seen: set[str] = set()

    for value in raw_headers:
        normalized = normalize_header(value)
        key = HEADER_ALIASES.get(normalized)
        if key and key in seen:
            raise SpreadsheetValidationError(
                f"La columna {key} aparece más de una vez."
            )
        if key:
            seen.add(key)
        mapped.append(key)

    missing = sorted(REQUIRED_COLUMNS - seen)
    if missing:
        raise SpreadsheetValidationError(
            "Faltan columnas obligatorias: " + ", ".join(missing)
        )
    return mapped


def normalize_header(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\ufffd", "")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = re.sub(r"[^0-9a-zA-Z]+", " ", text).strip().casefold()
    return re.sub(r"\s+", " ", text)


def _normalize_row(
    row_number: int,
    raw_headers: list[Any],
    header_keys: list[str | None],
    raw_values: list[Any],
) -> NormalizedDocumentRow:
    source_values: dict[str, Any] = {}
    mapped_values: dict[str, Any] = {}
    normalization_errors: list[str] = []

    for index, value in enumerate(raw_values):
        header = str(raw_headers[index] or f"Columna{index + 1}")
        source_values[header] = _json_value(value)
        key = header_keys[index]
        if key:
            mapped_values[key] = value

    values: dict[str, Any] = {
        "tipo_documento": _text(mapped_values.get("tipo_documento")),
        "folio": _identifier_text(mapped_values.get("folio")),
        "prefijo": _text(mapped_values.get("prefijo")),
        "divisa": _upper_text(mapped_values.get("divisa")),
        "forma_pago": _identifier_text(mapped_values.get("forma_pago")),
        "medio_pago": _identifier_text(mapped_values.get("medio_pago")),
        "nit_emisor": _identifier_text(mapped_values.get("nit_emisor")),
        "nombre_emisor": _text(mapped_values.get("nombre_emisor")),
        "nit_receptor": _identifier_text(mapped_values.get("nit_receptor")),
        "nombre_receptor": _text(mapped_values.get("nombre_receptor")),
        "estado_dian": _text(mapped_values.get("estado_dian")),
        "grupo": _text(mapped_values.get("grupo")),
    }

    document_key = _lower_text(mapped_values.get("clave_documento"))
    values["clave_documento"] = document_key
    values["tipo_clave"] = _infer_key_type(values["tipo_documento"])

    values["fecha_emision"] = _parse_date(
        mapped_values.get("fecha_emision"), "Fecha Emisión", normalization_errors
    )
    values["fecha_recepcion"] = _parse_datetime(
        mapped_values.get("fecha_recepcion"),
        "Fecha Recepción",
        normalization_errors,
    )

    for key in DECIMAL_COLUMNS:
        values[key] = _parse_decimal(
            mapped_values.get(key), key, normalization_errors, default_zero=True
        )
    values["total"] = _parse_decimal(
        mapped_values.get("total"), "total", normalization_errors, default_zero=False
    )

    if normalization_errors:
        source_values["_erroresNormalizacion"] = normalization_errors

    original_json = json.dumps(
        source_values,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    row_hash = hash_canonical(source_values)
    content_hash = "" if normalization_errors else hash_canonical(values)

    return NormalizedDocumentRow(
        row_number=row_number,
        document_key=document_key,
        key_type=values["tipo_clave"],
        values=values,
        row_hash=row_hash,
        content_hash=content_hash,
        original_json=original_json,
    )


def _infer_key_type(document_type: str | None) -> str:
    normalized = normalize_header(document_type)
    cude_markers = (
        "application response",
        "acuse",
        "aceptacion",
        "reclamo",
        "evento",
        "recibo",
    )
    return "CUDE" if any(marker in normalized for marker in cude_markers) else "CUFE"


def _parse_date(value: Any, label: str, errors: list[str]) -> date | None:
    if value is None or _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for pattern in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    errors.append(f"{label} no tiene un formato reconocido")
    return None


def _parse_datetime(
    value: Any, label: str, errors: list[str]
) -> datetime | None:
    if value is None or _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.replace(microsecond=0)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    text = str(value).strip()
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(
            tzinfo=None, microsecond=0
        )
    except ValueError:
        pass
    for pattern in ("%d-%m-%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    errors.append(f"{label} no tiene un formato reconocido")
    return None


def _parse_decimal(
    value: Any,
    label: str,
    errors: list[str],
    default_zero: bool,
) -> Decimal | None:
    if value is None or _is_blank(value):
        return Decimal("0") if default_zero else None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return Decimal(str(value))

    text = str(value).strip().replace("$", "").replace(" ", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")

    try:
        return Decimal(text)
    except InvalidOperation:
        errors.append(f"{label} no es numérico")
        return None


def _text(value: Any) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip()
    return normalized or None


def _lower_text(value: Any) -> str | None:
    normalized = _text(value)
    return normalized.lower() if normalized else None


def _upper_text(value: Any) -> str | None:
    normalized = _text(value)
    return normalized.upper() if normalized else None


def _identifier_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, Decimal) and value == value.to_integral_value():
        return str(value.quantize(Decimal("1")))
    return _text(value)


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return value


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())

