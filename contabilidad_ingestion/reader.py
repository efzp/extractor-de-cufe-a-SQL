from __future__ import annotations

import json
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from dian_ingestion.errors import SpreadsheetValidationError

from .models import AccountingRow, AccountingWorkbook


HEADERS = {
    "FECHA": "Fecha",
    "DOCUMENTO": "Documento",
    "TIPODOC": "TipoDoc",
    "NUMDOC": "NumDoc",
    "CUENTA": "Cuenta",
    "NOMCUENTA": "NomCuenta",
    "CONCEPTO": "Concepto",
    "NATURALEZA": "Naturaleza",
    "CENTRO": "Centro",
    "CODIGOCTABANCARIA": "CodigoCtaBancaria",
    "CODIGOUSUARIO": "CodigoUsuario",
    "DEBITO": "Debito",
    "CREDITO": "Credito",
    "IDENTIDADTERCERO": "IdentidadTercero",
    "DOCFUENTE": "DocFuente",
    "FECHASISTEMA": "FechaSistema",
    "INDCONTABILIDAD": "IndContabilidad",
    "DV": "Dv",
    "NOMBRETERCERO": "NombreTercero",
    "CUENTABANCARIA": "CuentaBancaria",
    "NOMCENTRO": "NomCentro",
    "DESCRIPCIONCORTA": "DescripcionCorta",
    "DIRECCION": "Direccion",
    "TELEFONOS": "Telefonos",
    "NUMEROMOVIL": "NumeroMovil",
    "NOMCIUDAD": "NomCiudad",
    "CC": "CC",
}

REQUIRED = frozenset(
    {
        "Fecha", "Documento", "TipoDoc", "NumDoc", "Cuenta",
        "Naturaleza", "Debito", "Credito", "IdentidadTercero",
        "IndContabilidad",
    }
)


def _header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]", "", text.upper())


def _column_map(first_row: tuple) -> dict[int, str] | None:
    columns: dict[int, str] = {}
    found: set[str] = set()
    for index, cell in enumerate(first_row):
        key = HEADERS.get(_header(cell.value))
        if key:
            if key in found:
                raise SpreadsheetValidationError(
                    f"La columna contable {key} esta repetida."
                )
            columns[index] = key
            found.add(key)
    return columns if REQUIRED <= found else None


def _text(cell) -> str | None:
    value = cell.value
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        result = str(value)
    elif isinstance(value, float):
        result = format(Decimal(str(value)), "f")
    else:
        result = str(value).strip()
    if not result:
        return None
    number_format = getattr(cell, "number_format", "")
    if isinstance(value, (int, float)) and re.fullmatch(r"0{2,}", number_format):
        if Decimal(str(value)) == int(value):
            result = str(int(value)).zfill(len(number_format))
    return result


def _normalized_value(column: str, cell) -> str | None:
    value = _text(cell)
    if value is None:
        return None
    if column == "Fecha":
        if isinstance(cell.value, (datetime, date)):
            return cell.value.date().isoformat() if isinstance(cell.value, datetime) else value
        for pattern in ("%Y/%m/%d", "%Y-%m-%d"):
            try:
                return datetime.strptime(value, pattern).date().isoformat()
            except ValueError:
                pass
    elif column == "FechaSistema":
        if isinstance(cell.value, datetime):
            return cell.value.isoformat(timespec="milliseconds")
        try:
            parsed = datetime.fromisoformat(value)
            if parsed.tzinfo is None:
                return parsed.isoformat(timespec="milliseconds")
        except ValueError:
            pass
    return value


def read_accounting_workbook(
    content: bytes, preferred_sheet: str | None = None, max_rows: int = 50000
) -> AccountingWorkbook:
    if not content or max_rows <= 0:
        raise SpreadsheetValidationError("Archivo o limite de filas invalido.")

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise SpreadsheetValidationError("No fue posible abrir el XLSX contable.") from error

    try:
        if preferred_sheet is not None:
            if preferred_sheet not in workbook.sheetnames:
                raise SpreadsheetValidationError("La hoja indicada no existe en el XLSX.")
            candidates = [workbook[preferred_sheet]]
        else:
            candidates = workbook.worksheets

        selected = None
        header_row = 0
        columns = None
        for sheet in candidates:
            for cells in sheet.iter_rows(min_row=1, max_row=10):
                mapping = _column_map(cells)
                if mapping is not None:
                    selected = sheet
                    header_row = cells[0].row
                    columns = mapping
                    break
            if selected is not None:
                break
        if selected is None or columns is None:
            raise SpreadsheetValidationError(
                "No se encontro una hoja con las columnas contables obligatorias."
            )

        rows: list[AccountingRow] = []
        for cells in selected.iter_rows(min_row=header_row + 1):
            if not any(_text(cell) is not None for cell in cells):
                continue
            if len(rows) >= max_rows:
                raise SpreadsheetValidationError(
                    f"El XLSX supera el limite de {max_rows} movimientos."
                )
            payload = {
                name: _normalized_value(name, cells[index])
                if index < len(cells) else None
                for index, name in columns.items()
            }
            rows.append(
                AccountingRow(
                    row_number=cells[0].row,
                    json_text=json.dumps(
                        payload, ensure_ascii=False, separators=(",", ":")
                    ),
                )
            )
        if not rows:
            raise SpreadsheetValidationError("La hoja contable no tiene movimientos.")
        return AccountingWorkbook(sheet_name=selected.title, rows=tuple(rows))
    finally:
        workbook.close()
